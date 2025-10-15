#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Script to analyze sample Excel files and extract their structure
"""

import pandas as pd
from openpyxl import load_workbook
import json

def analyze_excel_structure(file_path):
    """Analyze Excel file structure and return detailed information"""
    print(f"\n{'='*80}")
    print(f"Analyzing: {file_path}")
    print(f"{'='*80}")

    try:
        # Load workbook
        wb = load_workbook(file_path, data_only=True)

        result = {
            'file': file_path,
            'sheets': {}
        }

        # Analyze each sheet
        for sheet_name in wb.sheetnames:
            print(f"\n--- Sheet: {sheet_name} ---")
            ws = wb[sheet_name]

            # Get headers (first non-empty row)
            headers = []
            header_row = 1

            for row_idx in range(1, min(20, ws.max_row + 1)):
                row_values = []
                for cell in ws[row_idx]:
                    if cell.value is not None:
                        row_values.append(str(cell.value).strip())
                    else:
                        row_values.append('')

                # Check if this looks like a header row
                if len([v for v in row_values if v]) >= 5:  # At least 5 non-empty cells
                    headers = row_values
                    header_row = row_idx
                    break

            print(f"Header row: {header_row}")
            print(f"Headers ({len([h for h in headers if h])} columns):")
            for idx, header in enumerate(headers[:30], 1):  # Show first 30 headers
                if header:
                    print(f"  Col {idx}: {header}")

            # Get sample data (first 3 data rows)
            print(f"\nSample data (first 3 rows):")
            data_rows = []
            for row_idx in range(header_row + 1, min(header_row + 4, ws.max_row + 1)):
                row_data = {}
                for col_idx, cell in enumerate(ws[row_idx], 1):
                    if col_idx <= len(headers) and headers[col_idx-1]:
                        value = cell.value
                        if value is not None:
                            row_data[headers[col_idx-1]] = str(value)[:50]  # Limit to 50 chars

                if row_data:
                    data_rows.append(row_data)
                    print(f"  Row {row_idx}: {json.dumps(row_data, ensure_ascii=False, indent=4)}")

            # Store sheet info
            result['sheets'][sheet_name] = {
                'header_row': header_row,
                'headers': [h for h in headers if h],
                'total_rows': ws.max_row,
                'total_cols': ws.max_column,
                'sample_data': data_rows
            }

        return result

    except Exception as e:
        print(f"Error analyzing {file_path}: {e}")
        import traceback
        traceback.print_exc()
        return None

def analyze_mappings(file_path):
    """Analyze mappings.xlsx to extract status mappings"""
    print(f"\n{'='*80}")
    print(f"Analyzing mappings: {file_path}")
    print(f"{'='*80}")

    try:
        # Try to read as Excel
        wb = load_workbook(file_path, data_only=True)

        mappings = {}

        for sheet_name in wb.sheetnames:
            print(f"\n--- Sheet: {sheet_name} ---")
            ws = wb[sheet_name]

            # Read all data
            data = []
            for row in ws.iter_rows(min_row=1, values_only=True):
                row_data = [str(cell).strip() if cell is not None else '' for cell in row]
                if any(row_data):  # Skip empty rows
                    data.append(row_data)

            # Display the mapping data
            for row_idx, row in enumerate(data[:20], 1):  # Show first 20 rows
                print(f"  Row {row_idx}: {row}")

            mappings[sheet_name] = data

        return mappings

    except Exception as e:
        print(f"Error analyzing mappings {file_path}: {e}")
        import traceback
        traceback.print_exc()
        return None

if __name__ == '__main__':
    base_path = r'C:\Users\Administrator\Desktop\TrackerBuilder'

    # Analyze sample1.xlsx
    sample1_info = analyze_excel_structure(f'{base_path}\\sample1.xlsx')

    # Analyze sample2.xlsx
    sample2_info = analyze_excel_structure(f'{base_path}\\sample2.xlsx')

    # Analyze mappings.xlsx
    mappings_info = analyze_mappings(f'{base_path}\\mappings.xlsx')

    # Save results to JSON
    results = {
        'sample1': sample1_info,
        'sample2': sample2_info,
        'mappings': mappings_info
    }

    with open('excel_analysis_results.json', 'w', encoding='utf-8') as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    print(f"\n{'='*80}")
    print("Analysis complete! Results saved to excel_analysis_results.json")
    print(f"{'='*80}")
