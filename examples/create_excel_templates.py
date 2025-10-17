"""
Create Example Excel Templates for Template Filling Feature

This script generates example Excel templates that can be used with the
/api/v1/excel/fill-template endpoint.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime


def create_progress_report_template():
    """Create a project progress report template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "项目进度报表"

    # Set column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 20

    # Title row (merged)
    ws.merge_cells('A1:H1')
    title_cell = ws['A1']
    title_cell.value = "AK/云原生转型项目进度报表"
    title_cell.font = Font(name='微软雅黑', size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Timestamp row (merged)
    ws.merge_cells('A2:H2')
    timestamp_cell = ws['A2']
    timestamp_cell.value = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    timestamp_cell.font = Font(name='微软雅黑', size=10, italic=True)
    timestamp_cell.alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 20

    # Header row
    headers = ['L2 ID', '应用名称', '转型目标', '当前状态', '进度%', '开发团队', '负责人', '备注']
    header_fill = PatternFill(start_color='B4C7E7', end_color='B4C7E7', fill_type='solid')
    header_font = Font(name='微软雅黑', size=11, bold=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[3].height = 25

    # Add a few empty data rows with borders
    for row in range(4, 7):
        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            cell.border = border

    return wb


def create_delayed_projects_template():
    """Create a delayed projects analysis template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "延期项目分析"

    # Set column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 12
    ws.column_dimensions['I'].width = 20

    # Title row (merged) - Red background for attention
    ws.merge_cells('A1:I1')
    title_cell = ws['A1']
    title_cell.value = "延期项目分析报表"
    title_cell.font = Font(name='微软雅黑', size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Timestamp row
    ws.merge_cells('A2:I2')
    timestamp_cell = ws['A2']
    timestamp_cell.value = f"统计时间: {datetime.now().strftime('%Y-%m-%d')}"
    timestamp_cell.font = Font(name='微软雅黑', size=10)
    timestamp_cell.alignment = Alignment(horizontal='center')

    # Header row
    headers = ['应用ID', '应用名称', '负责团队', '计划上线日期', '实际上线日期',
               '延期天数', '当前进度', '负责人', '延期原因']
    header_fill = PatternFill(start_color='F4B084', end_color='F4B084', fill_type='solid')
    header_font = Font(name='微软雅黑', size=11, bold=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[3].height = 25

    # Add empty data rows
    for row in range(4, 7):
        for col in range(1, 10):
            cell = ws.cell(row=row, column=col)
            cell.border = border

    return wb


def create_team_statistics_template():
    """Create a team statistics template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "团队统计"

    # Set column widths
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15

    # Title row
    ws.merge_cells('A1:G1')
    title_cell = ws['A1']
    title_cell.value = "团队转型项目统计表"
    title_cell.font = Font(name='微软雅黑', size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Header row
    headers = ['团队名称', '项目总数', '已完成', '进行中', '未开始', '平均进度', '延期项目数']
    header_fill = PatternFill(start_color='A9D08E', end_color='A9D08E', fill_type='solid')
    header_font = Font(name='微软雅黑', size=11, bold=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[2].height = 25

    # Add empty data rows
    for row in range(3, 6):
        for col in range(1, 8):
            cell = ws.cell(row=row, column=col)
            cell.border = border

    return wb


def create_simple_application_list_template():
    """Create a simple application list template for basic testing."""
    wb = Workbook()
    ws = wb.active
    ws.title = "应用列表"

    # Set column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 15

    # Header row with simple styling
    headers = ['L2 ID', '应用名称', '状态', '进度', '团队']
    header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    header_font = Font(name='微软雅黑', size=11, bold=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[1].height = 20

    return wb


def main():
    """Generate all example templates."""
    # Create examples directory if it doesn't exist
    output_dir = os.path.join(os.path.dirname(__file__), 'excel_templates')
    os.makedirs(output_dir, exist_ok=True)

    templates = [
        ('template_progress_report.xlsx', create_progress_report_template),
        ('template_delayed_projects.xlsx', create_delayed_projects_template),
        ('template_team_statistics.xlsx', create_team_statistics_template),
        ('template_simple_list.xlsx', create_simple_application_list_template),
    ]

    print("Creating example Excel templates...")
    for filename, create_func in templates:
        wb = create_func()
        filepath = os.path.join(output_dir, filename)
        wb.save(filepath)
        print(f"[OK] Created: {filepath}")

    print(f"\n[SUCCESS] All templates created successfully in: {output_dir}")
    print("\nUsage:")
    print("  Upload these templates to the /api/v1/excel/fill-template endpoint")
    print("  The system will automatically fill them with data from the database")


if __name__ == '__main__':
    main()
